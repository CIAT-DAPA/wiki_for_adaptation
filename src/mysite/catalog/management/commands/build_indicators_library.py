"""
Parse the Indicators planning spreadsheet into a JSON file the site can render.

The "extended indicators library" is an *informational* page listing every
planned indicator/metric — content that is not (yet) loaded as live wiki pages.
Rather than depend on the spreadsheet at runtime, we parse it once into
``catalog/data/indicators_library.json`` (committed to the repo) and the view
reads that. Re-run this command whenever the spreadsheet changes:

    python manage.py build_indicators_library --source path/to/Indicators.xlsx

Only the first sheet ("Indicators") is used. Dimension / Subdimension /
Indicator cells are merged in the spreadsheet (value only on the first row of
each group), so they are forward-filled down.
"""
from __future__ import annotations

import json
from collections import OrderedDict
from datetime import date
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError

# Column order in the "Indicators" sheet.
COL_DIMENSION = 0        # "Dimension (aka Purpose)"
COL_SUBDIMENSION = 1     # "Subdimension (aka Theme)"
COL_INDICATOR = 2        # "Indicator"
COL_METRIC = 3           # "Metric"
COL_METRIC_TYPE = 6      # "Metric type"
COL_PHASE = 7            # "Publishing phase"
COL_DEFINITION = 8       # "Metric Definition"
COL_REFERENCES = 9       # "References"

OUTPUT_PATH = Path(__file__).resolve().parents[2] / "data" / "indicators_library.json"


def _clean(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


class Command(BaseCommand):
    help = "Parse the Indicators spreadsheet into catalog/data/indicators_library.json"

    def add_arguments(self, parser):
        parser.add_argument("--source", required=True, help="Path to Indicators.xlsx")
        parser.add_argument("--sheet", default="Indicators", help="Sheet name (default: Indicators)")

    def handle(self, *args, **options):
        try:
            import openpyxl
        except ImportError:
            raise CommandError("openpyxl is required. Add it to requirements and pip install.")

        source = Path(options["source"])
        if not source.exists():
            raise CommandError(f"Source file not found: {source}")

        wb = openpyxl.load_workbook(source, data_only=True, read_only=True)
        sheet_name = options["sheet"]
        if sheet_name not in wb.sheetnames:
            raise CommandError(f"Sheet {sheet_name!r} not found. Sheets: {wb.sheetnames}")
        ws = wb[sheet_name]

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise CommandError("Spreadsheet is empty.")

        data_rows = rows[1:]  # skip header

        # Forward-fill the merged grouping columns.
        last = {}
        filled = []
        for raw in data_rows:
            raw = list(raw) + [None] * (10 - len(raw))  # pad short rows
            for col in (COL_DIMENSION, COL_SUBDIMENSION, COL_INDICATOR):
                if _clean(raw[col]):
                    last[col] = _clean(raw[col])
                else:
                    raw[col] = last.get(col, "")
            if any(_clean(c) for c in raw):
                filled.append(raw)

        # Build the nested structure, preserving spreadsheet order.
        dims: "OrderedDict[str, OrderedDict]" = OrderedDict()
        metric_count = 0
        indicator_names = set()
        for r in filled:
            dim = _clean(r[COL_DIMENSION]) or "Uncategorised"
            sub = _clean(r[COL_SUBDIMENSION]) or "Uncategorised"
            ind = _clean(r[COL_INDICATOR])
            metric_name = _clean(r[COL_METRIC])

            sub_map = dims.setdefault(dim, OrderedDict())
            ind_map = sub_map.setdefault(sub, OrderedDict())
            if not ind:
                continue
            indicator_names.add((dim, sub, ind))
            metrics = ind_map.setdefault(ind, [])
            if metric_name:
                metrics.append({
                    "name": metric_name,
                    "type": _clean(r[COL_METRIC_TYPE]),
                    "phase": _clean(r[COL_PHASE]),
                    "definition": _clean(r[COL_DEFINITION]),
                    "references": _clean(r[COL_REFERENCES]),
                })
                metric_count += 1

        dimensions = []
        for dim, sub_map in dims.items():
            subdimensions = []
            for sub, ind_map in sub_map.items():
                indicators = [{"name": name, "metrics": metrics} for name, metrics in ind_map.items()]
                subdimensions.append({"name": sub, "indicators": indicators})
            dimensions.append({"name": dim, "subdimensions": subdimensions})

        payload = {
            "source": f"{source.name} :: {sheet_name}",
            "generated_at": date.today().isoformat(),
            "counts": {
                "dimensions": len(dimensions),
                "subdimensions": sum(len(d["subdimensions"]) for d in dimensions),
                "indicators": len(indicator_names),
                "metrics": metric_count,
            },
            "dimensions": dimensions,
        }

        OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
        OUTPUT_PATH.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")

        c = payload["counts"]
        self.stdout.write(self.style.SUCCESS(
            f"Wrote {OUTPUT_PATH.name}: {c['dimensions']} dimensions, "
            f"{c['subdimensions']} subdimensions, {c['indicators']} indicators, "
            f"{c['metrics']} metrics."
        ))
