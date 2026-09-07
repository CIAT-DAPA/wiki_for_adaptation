"""
Reusable bulk-import logic for the SOP masterfile.

Shared by the ``import_sops`` management command (console) and the Wagtail admin
upload view, so both behave identically. ``run_import`` accepts a file path or a
file-like object (e.g. an uploaded file) and returns a structured report.

Each row of the "SOPs" sheet is one SOP, carrying its Indicator and Metric
context. Pages are matched by a stable ``external_id`` so re-running updates
instead of duplicating.
"""
from __future__ import annotations

import re

from django.contrib.auth import get_user_model
from django.db import transaction
from django.utils.html import escape
from django.utils.text import slugify
from wagtail.coreutils import find_available_slug

MAX_METHODS = 5
_URL_RE = re.compile(r"(https?://[^\s;]+)")


class SopImportError(Exception):
    """Raised for problems that abort the whole import (bad file/sheet/setup)."""


# --------------------------------------------------------------------------- #
# Text -> HTML helpers
# --------------------------------------------------------------------------- #

def _txt(value) -> str:
    return "" if value is None else str(value).strip()


def _strip_sop_prefix(title: str) -> str:
    # The metric page template already renders "Standard Operating Procedure - "
    # before the SOP title, so drop the prefix if the spreadsheet includes it.
    return re.sub(r"^\s*standard operating procedure\s*[-–—:]\s*", "", title, flags=re.I).strip()


def _linkify(escaped: str) -> str:
    return _URL_RE.sub(lambda m: f'<a href="{m.group(1)}">{m.group(1)}</a>', escaped)


def _paragraphs(value) -> str:
    text = _txt(value)
    if not text:
        return ""
    parts = [p.strip() for p in text.split("\n") if p.strip()]
    return "".join(f"<p>{escape(p)}</p>" for p in parts)


def _list_html(value, tag: str = "ul", linkify: bool = False) -> str:
    text = _txt(value)
    if not text:
        return ""
    items = [p.strip() for p in text.split("\n") if p.strip()]
    if len(items) <= 1:
        items = [p.strip() for p in text.split(";") if p.strip()]
    if not items:
        return ""
    if len(items) == 1:
        cell = _linkify(escape(items[0])) if linkify else escape(items[0])
        return f"<p>{cell}</p>"
    lis = []
    for i in items:
        cell = _linkify(escape(i)) if linkify else escape(i)
        lis.append(f"<li>{cell}</li>")
    return f"<{tag}>{''.join(lis)}</{tag}>"


# --------------------------------------------------------------------------- #
# Field builders
# --------------------------------------------------------------------------- #

def _indicator_fields(row):
    return {
        "title": _txt(row.get("indicator_name"))[:255],
        "description": _paragraphs(row.get("indicator_description")),
        "dimension": _txt(row.get("indicator_properties"))[:150],
        "external_id": _txt(row.get("indicator_id")),
    }


def _metric_fields(row):
    tf = _txt(row.get("metric_tracking_function")).lower()
    return {
        "title": _txt(row.get("metric_name"))[:255],
        "description": _paragraphs(row.get("metric_description")),
        "purpose": _paragraphs(row.get("metric_purpose")),
        "tracks_vulnerability": "vulnerab" in tf,
        "tracks_intervention_impacts": ("intervention" in tf) or ("impact" in tf),
        "adaptation_tracking_function": _paragraphs(row.get("metric_tracking_function")),
    }


def _sop_fields(row):
    return {
        "title": _strip_sop_prefix(_txt(row.get("sop_title")))[:255],
        "external_id": _txt(row.get("sop_id")),
        "definition": _paragraphs(row.get("sop_definition")),
        "data_sources": _list_html(row.get("sop_data_sources")),
        "units": _list_html(row.get("sop_units")),
        "frequency": _paragraphs(row.get("sop_frequency")),
        "geographic_scale": _paragraphs(row.get("sop_scale")),
        "technical_capacity": _paragraphs(row.get("sop_capacity")),
        "activities_and_steps": _list_html(row.get("sop_activities_steps"), "ol"),
        "options_enhancing_robustness": _list_html(row.get("sop_robustness")),
        "options_reducing_costs": _list_html(row.get("sop_cost")),
        "available_tools_and_code": _list_html(row.get("sop_tools"), linkify=True),
        "references": _list_html(row.get("sop_references"), linkify=True),
        "flagship_method_status": _paragraphs(row.get("sop_flagship_status")),
        "visual_content": _paragraphs(row.get("sop_visual_content")),
        "example_applications": _list_html(row.get("sop_example_applications")),
        "unfccc_alignment": _paragraphs(row.get("sop_unfccc_alignment")),
    }


def _method_field_sets(row):
    for n in range(1, MAX_METHODS + 1):
        desc = _txt(row.get(f"method{n}_description"))
        name = _txt(row.get(f"method{n}_name"))
        if not (desc or name):
            continue
        yield {
            "title": (name or f"Method option {n}")[:255],
            "description": _paragraphs(row.get(f"method{n}_description")),
            "resolution": _paragraphs(row.get(f"method{n}_resolution")),
            "advantages": _list_html(row.get(f"method{n}_advantages")),
            "limitations": _list_html(row.get(f"method{n}_limitations")),
            "use_case": _paragraphs(row.get(f"method{n}_use_case")),
            "resources": _list_html(row.get(f"method{n}_resources"), linkify=True),
        }


# --------------------------------------------------------------------------- #
# Import engine
# --------------------------------------------------------------------------- #

def load_records(source, sheet: str = "SOPs"):
    """Read the sheet into a list of {column: value} dicts."""
    try:
        import openpyxl
    except ImportError as exc:
        raise SopImportError("openpyxl is required (pip install -r requirements.txt).") from exc

    try:
        wb = openpyxl.load_workbook(source, data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise SopImportError(f"Could not read the spreadsheet: {exc}") from exc

    if sheet not in wb.sheetnames:
        raise SopImportError(f"Sheet {sheet!r} not found. Sheets: {wb.sheetnames}")
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        raise SopImportError("Spreadsheet has no data rows.")
    header = [_txt(c) for c in rows[0]]
    return [dict(zip(header, r)) for r in rows[1:] if any(c is not None for c in r)]


class _Importer:
    def __init__(self, dry_run, publish, owner):
        self.dry_run = dry_run
        self.publish = publish
        self.owner = owner

    def process_row(self, row, home):
        from catalog.models import IndicatorPage, MetricPage, MethodPage, SOPPage

        ind_fields = _indicator_fields(row)
        metric_fields = _metric_fields(row)
        sop_fields = _sop_fields(row)

        if not ind_fields["title"]:
            raise ValueError("missing indicator_name")
        if not metric_fields["title"]:
            raise ValueError("missing metric_name")
        if not sop_fields["title"]:
            raise ValueError("missing sop_title")

        indicator = None
        if ind_fields["external_id"]:
            indicator = IndicatorPage.objects.filter(external_id=ind_fields["external_id"]).first()
        if indicator is None:
            indicator = IndicatorPage.objects.child_of(home).filter(title=ind_fields["title"]).first()
        if indicator is None:
            indicator = self._create_page(IndicatorPage, home, ind_fields)
            action = "created"
        else:
            self._update_page(indicator, ind_fields)
            action = "updated"

        sop = None
        if sop_fields["external_id"]:
            sop = SOPPage.objects.filter(external_id=sop_fields["external_id"]).first()

        if sop is not None:
            metric = sop.get_parent().specific
            self._update_page(metric, metric_fields)
            self._update_page(sop, sop_fields)
        else:
            metric = MetricPage.objects.child_of(indicator).filter(title=metric_fields["title"]).first()
            if metric is None:
                metric = self._create_page(MetricPage, indicator, metric_fields)
            else:
                self._update_page(metric, metric_fields)
            self._create_page(SOPPage, metric, sop_fields)
            action = "created"

        for m_fields in _method_field_sets(row):
            existing = MethodPage.objects.child_of(metric).filter(title=m_fields["title"]).first()
            if existing is None:
                self._create_page(MethodPage, metric, m_fields)
            else:
                self._update_page(existing, m_fields)

        return action

    def _create_page(self, model, parent, fields):
        page = model(**fields)
        page.slug = find_available_slug(parent, slugify(fields["title"]) or "item")
        page.live = self.publish
        page.owner = self.owner
        parent.add_child(instance=page)
        revision = page.save_revision(user=self.owner)
        if self.publish:
            revision.publish()
        return page

    def _update_page(self, page, fields):
        for key, value in fields.items():
            if key == "external_id" and not value:
                continue
            setattr(page, key, value)
        if page.owner_id is None:
            page.owner = self.owner
        revision = page.save_revision(user=self.owner)
        if self.publish:
            revision.publish()


def run_import(source, *, sheet="SOPs", dry_run=False, publish=False, owner=None):
    """Import the spreadsheet and return a report dict.

    report = {mode, total, created, updated, errors, rows:[{action, sop_id, title, error?}]}
    On a dry run every row is processed inside a transaction that is rolled back.
    """
    from home.models import HomePage

    records = load_records(source, sheet)

    home = HomePage.objects.first()
    if home is None:
        raise SopImportError("No HomePage found to attach indicators to.")

    if owner is None:
        User = get_user_model()
        owner = (
            User.objects.filter(is_superuser=True).order_by("id").first()
            or User.objects.order_by("id").first()
        )

    importer = _Importer(dry_run, publish, owner)
    report = {
        "mode": "dry-run" if dry_run else ("publish" if publish else "draft"),
        "total": len(records),
        "created": 0,
        "updated": 0,
        "errors": 0,
        "rows": [],
    }

    for i, row in enumerate(records, start=1):
        sop_id = _txt(row.get("sop_id")) or f"row {i}"
        title = _txt(row.get("sop_title"))
        try:
            with transaction.atomic():
                action = importer.process_row(row, home)
                if dry_run:
                    transaction.set_rollback(True)
            report[action] += 1
            report["rows"].append({"action": action, "sop_id": sop_id, "title": title})
        except Exception as exc:  # noqa: BLE001 - collect and continue
            report["errors"] += 1
            report["rows"].append({"action": "error", "sop_id": sop_id, "title": title, "error": str(exc)})

    return report
